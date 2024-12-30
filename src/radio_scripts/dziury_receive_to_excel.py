import sys
import os
# Add the parent directory to the system path
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import radio_utils.radio_utils as radio_utils
from time import time
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook  # Import to manage existing Excel files




def send_data(serial_conn, data):
    return serial_conn.write(data)

def receive_data(serial_conn) -> str:
    """Receives data from serial connection, expecting 50-byte messages ending in 'S'."""
    buff = ""
    while True:
        byte = serial_conn.read()
        if byte == b'':
            continue
        if byte == b'G':  # Message starts with 'G'
            buff = ""
        try:  # Handle decoding errors
            buff += byte.decode()
        except:
            pass
        if len(buff) == 50 and buff.endswith("S"):  # Full 50-byte message received
            return buff
        
def append_to_excel(filename, data):
    """
    Appends data to an existing Excel file or creates a new one if it doesn't exist.

    Args:
        filename (str): The path to the Excel file.
        data (dict): The data to append as a dictionary of lists.
    """
    df = pd.DataFrame(data)

    try:
        # Try to open an existing workbook
        with pd.ExcelWriter(filename, mode='a', if_sheet_exists='overlay', engine='openpyxl') as writer:
            # Find the current max row in the existing sheet
            book = writer.book
            if 'Sheet1' in book.sheetnames:
                sheet = book['Sheet1']
                start_row = sheet.max_row
            else:
                start_row = 0

            # Append the DataFrame starting from the next row
            df.to_excel(writer, index=False, startrow=start_row, header=False)
    except FileNotFoundError:
        # If file doesn't exist, create a new one
        df.to_excel(filename, index=False)

def main():
    if len(sys.argv) > 1:
        selected_port = sys.argv[1]
        min_freq = int(sys.argv[2])
        max_freq = int(sys.argv[3])
        reading_period = int(sys.argv[4])  # Message read duration in seconds
        sending_freq = int(sys.argv[5])  # Messages sent per second
        air_speed = int(sys.argv[6])
    else:
        selected_port = "COM5"
        reading_period = 7

    detected_baud = 57600

    last_seqNum = -1
    last_send_timestamp = 0
    start_pelne = 0
    good_msgs = 0

    send_timestamps = []
    recv_timestamps = []
    lost_timestamps = []
    wrong_timestamps = []
    lost_y = []
    send_y = []
    recv_y = []
    wrong_y = []
    wrong_y_val = 0
    total_bytes_received = 0
    first_wrong_found = False

    dziury = []
    pelne = []

    try:
        with radio_utils.serial.Serial(selected_port, detected_baud, timeout=0.000001) as ser:
            ser.flushInput()
            ser.flushOutput()
            start_time = time()

            while time() - start_time < reading_period:
                received_data = receive_data(ser)
                time_ms = int(time() * 1000)

                if received_data[0:2] != "G " or not received_data.endswith("S"):
                    # print("Incorrect message format:", received_data)
                    if wrong_y_val != 0:
                        wrong_timestamps.append(time_ms)
                        wrong_y_val += 1
                        wrong_y.append(wrong_y_val)
                    continue

                fields = received_data.split(" ")
                if len(fields) < 5 or '' in fields:
                    # print("Incorrect number of fields:", received_data)
                    if wrong_y_val != 0:
                        wrong_timestamps.append(time_ms)
                        wrong_y_val += 1
                        wrong_y.append(wrong_y_val)
                    continue

                seqNum = int(fields[1])
                sendTS = int(fields[2])
                checksum = seqNum + sendTS

                if checksum != int(fields[3]):
                    # print("Incorrect checksum:", received_data)
                    if wrong_y_val != 0:
                        wrong_timestamps.append(time_ms)
                        wrong_y_val += 1
                        wrong_y.append(wrong_y_val)
                    continue

                if start_pelne == 0:
                    start_pelne = sendTS
                if last_send_timestamp == 0:
                    last_send_timestamp = sendTS
                if last_seqNum == -1:
                    last_seqNum = seqNum - 1
                    wrong_y_val = seqNum - 1

                send_timestamps.append(sendTS)
                recv_timestamps.append(int(time() * 1000))
                recv_y.append(seqNum)
                send_y.append(seqNum)

                total_bytes_received += len(received_data.encode('utf-8'))
                


                # print(fields)

                good_msgs += 1
                if (last_seqNum + 1 != seqNum):
                    print(f'{last_seqNum} | {seqNum}')
                    if not first_wrong_found: first_wrong_found = True

                    num_of_lost_messages = seqNum - last_seqNum - 1
                    dziura_ts = sendTS - last_send_timestamp
                    dziury.append(dziura_ts)
                    delta_ts = dziura_ts / (num_of_lost_messages + 1)
                    # print("Last ts:", last_send_timestamp % 10000, "now ts", sendTS % 10000, "last_seq:", last_seqNum, "now_seq:", seqNum, "lost_msgs:", num_of_lost_messages, "delta_ts:", delta_ts)

                    for i in range(num_of_lost_messages):
                        send_timestamps.append(last_send_timestamp + (i + 1) * int(delta_ts))
                        lost_timestamps.append(last_send_timestamp + (i + 1) * int(delta_ts) + 3)
                        send_y.append(last_seqNum + i + 1)
                        lost_y.append(last_seqNum + i + 1)
                    
                    good_msgs = 0
                    pelne.append(last_send_timestamp - start_pelne)
                    start_pelne = sendTS
                last_seqNum = seqNum
                wrong_y_val = seqNum
                last_send_timestamp = sendTS

            if len(dziury) > 0:
                dziury_avg = sum(dziury) / len(dziury)
                pelne_avg = sum(pelne) / len(pelne)
                # print("Dziury:", dziury, "avg:", dziury_avg)
                # print("Pelne:", pelne, "avg:", pelne_avg)

            bandwidth = max_freq-min_freq
            sending_freq = 50*sending_freq
            output_filename = 'radio_test_results.xlsx'
            data = {
                'Messages sent': [len(send_timestamps)],
                'Messages lost': [len(lost_timestamps)],
                'Messages % lost': [(len(lost_timestamps) / len(send_timestamps))*100 if len(send_timestamps) > 0 else 0],
                'Bandwidth (Hz)': [abs(bandwidth)],
                'Sending speed (B/s)': [sending_freq],
                'Airspeed (B/s)': [air_speed*1000/8],
                'Baud rate / 10 (B/s)': [detected_baud / 10]
            }

            append_to_excel(output_filename, data)
            print(f'Data appended to {output_filename}')

            # Save the plot with a filename based on parameters
            # plot_filename = f'{air_speed}kbps_{reading_period}s_{sending_frequency}Hz_minf_{min_freq}_maxf_{max_freq}.png'

    except radio_utils.serial.SerialException as e:
        print(f'Error: {e}')
    except Exception as e:
        print(f'Unexpected error: {e}')

if __name__ == '__main__':
    main()